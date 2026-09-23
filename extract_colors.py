import json
import re
import openpyxl

SRC = r"C:\Users\user\Downloads\Cleaning Schedule 2026.xlsx"
DAYS = ['Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
LAST_ROW = {
    'Sunday': 80, 'Monday': 91, 'Tuesday': 91, 'Wednesday': 91,
    'Thursday': 91, 'Friday': 51, 'Saturday': 88,
}


def clean_staff(name):
    if not name:
        return None
    return re.sub(r'\s*-\s*.*$', '', name).strip()


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    colors = []

    for day in DAYS:
        ws = wb[day]
        staff_headers = [ws.cell(row=2, column=c).value for c in range(2, 8)]
        staff_names = [clean_staff(h) for h in staff_headers]

        for r in range(3, LAST_ROW[day] + 1):
            time_val = ws.cell(row=r, column=1).value
            if time_val is None:
                continue
            for ci, staff in enumerate(staff_names, start=2):
                cell = ws.cell(row=r, column=ci)
                task = cell.value
                if task is None or staff is None:
                    continue
                task_str = str(task).strip()
                if task_str in ('', ',', '-'):
                    continue
                fill = cell.fill
                rgb = fill.fgColor.rgb if fill and fill.fgColor else None
                if not isinstance(rgb, str) or rgb in ('00000000',):
                    rgb = None
                colors.append(rgb)

    with open("task_colors.json", "w") as f:
        json.dump(colors, f)
    print(f"Total colors extracted: {len(colors)}")

    # sanity check against existing schedule_normalized.json ordering
    tasks = json.load(open("schedule_normalized.json"))
    print(f"schedule_normalized.json length: {len(tasks)}")
    assert len(tasks) == len(colors), "LENGTH MISMATCH - order may not align!"
    print("Lengths match.")


if __name__ == "__main__":
    main()
