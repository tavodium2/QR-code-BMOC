import json
import re
import openpyxl
from collections import Counter

SRC = r"C:\Users\user\Downloads\Cleaning Schedule 2026.xlsx"
DAYS = ['Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
LAST_ROW = {
    'Sunday': 80, 'Monday': 91, 'Tuesday': 91, 'Wednesday': 91,
    'Thursday': 91, 'Friday': 51, 'Saturday': 88,
}


def clean_staff(name):
    if not name:
        return None
    return re.sub(r'\s*-\s*.*$', '', name).strip()


def normalize_task(t):
    return re.sub(r'\s+', ' ', t.strip().lower()).strip(',').strip()


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    all_tasks = []

    for day in DAYS:
        ws = wb[day]
        staff_headers = [ws.cell(row=2, column=c).value for c in range(2, 8)]
        staff_names = [clean_staff(h) for h in staff_headers]

        day_tasks = []
        for r in range(3, LAST_ROW[day] + 1):
            time_val = ws.cell(row=r, column=1).value
            if time_val is None:
                continue
            time_str = time_val.strftime('%H:%M') if hasattr(time_val, 'strftime') else str(time_val)
            for ci, staff in enumerate(staff_names, start=2):
                task = ws.cell(row=r, column=ci).value
                if task is None or staff is None:
                    continue
                task_str = str(task).strip()
                if task_str in ('', ',', '-'):
                    continue
                day_tasks.append({
                    'day': day, 'time': time_str, 'staff': staff, 'task': task_str,
                })

        # detect tasks that recur 2+ times in the same day (checkpoint-style tasks)
        norm_counts = Counter(normalize_task(t['task']) for t in day_tasks)
        for t in day_tasks:
            t['twice_daily'] = norm_counts[normalize_task(t['task'])] >= 2

        all_tasks.extend(day_tasks)

    with open('schedule.json', 'w') as f:
        json.dump(all_tasks, f, indent=2)

    print(f'Total task rows: {len(all_tasks)}')
    print(f'Twice-daily-flagged rows: {sum(1 for t in all_tasks if t["twice_daily"])}')
    by_day = Counter(t['day'] for t in all_tasks)
    for d in DAYS:
        print(f'  {d}: {by_day[d]}')


if __name__ == '__main__':
    main()
